#!/usr/bin/env python3
import argparse
import csv
import json
import os
import re
import sys
from collections import OrderedDict
from typing import Dict, List, Any, Optional, Tuple

from openpyxl import load_workbook

MASTER_SHEET = "Master"
REQUIRED_HEADERS = [
    "Pillar",
    "Business View Label",
    "Business View Name",
    "View Object",
    "View Object Attribute",
    "Business View Attribute",
    "Database Table",
    "Database Column",
]


def normalize(text: str) -> str:
    return re.sub(r"\s+", "", (text or "").strip()).lower()


def load_rows(workbook_path: str) -> List[Dict[str, Any]]:
    wb = load_workbook(workbook_path, read_only=True, data_only=True)
    if MASTER_SHEET not in wb.sheetnames:
        raise ValueError(f"Workbook does not contain required sheet: {MASTER_SHEET}")

    ws = wb[MASTER_SHEET]
    rows_iter = ws.iter_rows(values_only=True)
    headers = next(rows_iter, None)
    if not headers:
        raise ValueError("Master sheet is empty")

    header_positions = {str(h).strip(): idx for idx, h in enumerate(headers) if h is not None}
    missing = [h for h in REQUIRED_HEADERS if h not in header_positions]
    if missing:
        raise ValueError(f"Master sheet is missing required headers: {', '.join(missing)}")

    ordered_headers = REQUIRED_HEADERS
    records: List[Dict[str, Any]] = []
    for row in rows_iter:
        view_object = row[header_positions["View Object"]]
        if not view_object:
            continue
        record = {h: row[header_positions[h]] for h in ordered_headers}
        records.append(record)
    return records


def extract_version(path: str) -> str:
    filename = os.path.basename(path)
    match = re.search(r"_(\d{2}[A-Z])\.xlsx$", filename, re.IGNORECASE)
    return match.group(1).upper() if match else "unknown"


def build_index(records: List[Dict[str, Any]]) -> Dict[str, List[Dict[str, Any]]]:
    idx: Dict[str, List[Dict[str, Any]]] = {}
    for rec in records:
        key = normalize(str(rec["View Object"]))
        idx.setdefault(key, []).append(rec)
    return idx


def find_pvo(index: Dict[str, List[Dict[str, Any]]], pvo: str) -> List[Dict[str, Any]]:
    exact = index.get(normalize(pvo), [])
    if exact:
        return exact
    return []


def summarize_single(rows: List[Dict[str, Any]], pvo: str, version: str) -> OrderedDict:
    first = rows[0]
    return OrderedDict([
        ("version", version),
        ("pvo", pvo),
        ("business_view_name", first["Business View Name"]),
        ("pillar", first["Pillar"]),
        ("match_type", "exact"),
    ])


def summarize_mapping_with_columns(rows: List[Dict[str, Any]], pvo: str, version: str) -> Dict[str, Any]:
    first = rows[0]
    columns = []
    for row in rows:
        columns.append(OrderedDict([
            ("pvo", row["View Object"]),
            ("business_view_name", row["Business View Name"]),
            ("pvo_column", row["View Object Attribute"]),
            ("business_view_column", row["Business View Attribute"]),
            ("database_table", row["Database Table"]),
            ("database_column", row["Database Column"]),
            ("notes", "" if row["Business View Attribute"] != "Not in business view" else "not in business view"),
        ]))

    return OrderedDict([
        ("version", version),
        ("pvo", pvo),
        ("pillar", first["Pillar"]),
        ("business_view_label", first["Business View Label"]),
        ("business_view_name", first["Business View Name"]),
        ("match_type", "exact"),
        ("column_count", len(columns)),
        ("columns", columns),
    ])


def write_csv(mappings: List[Dict[str, Any]], out_path: str) -> None:
    headers = [
        "version",
        "pillar",
        "business_view_label",
        "business_view_name",
        "pvo",
        "pvo_column",
        "business_view_column",
        "database_table",
        "database_column",
        "notes",
    ]
    with open(out_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=headers)
        writer.writeheader()
        for mapping in mappings:
            base = {
                "version": mapping["version"],
                "pillar": mapping["pillar"],
                "business_view_label": mapping["business_view_label"],
                "business_view_name": mapping["business_view_name"],
                "pvo": mapping["pvo"],
            }
            for col in mapping["columns"]:
                row = dict(base)
                row.update({
                    "pvo_column": col["pvo_column"],
                    "business_view_column": col["business_view_column"],
                    "database_table": col["database_table"],
                    "database_column": col["database_column"],
                    "notes": col["notes"],
                })
                writer.writerow(row)


def resolve_csv_out_path(csv_out: str, version: str) -> str:
    """If csv_out points to a directory, generate a filename inside it."""
    if os.path.isdir(csv_out):
        return os.path.join(csv_out, f"bicc_to_boss_mapping_{version}.csv")
    return csv_out


def default_csv_out_path(version: str) -> str:
    return os.path.join("outputs", f"bicc_to_boss_mapping_{version}.csv")


def build_not_found_entry(pvo: str) -> OrderedDict:
    return OrderedDict([
        ("pvo", pvo),
        ("reason", "not present in workbook"),
    ])


def collect_results(
    pvos: List[str],
    index: Dict[str, List[Dict[str, Any]]],
) -> Tuple[List[Tuple[str, List[Dict[str, Any]]]], List[OrderedDict]]:
    found_rows: List[Tuple[str, List[Dict[str, Any]]]] = []
    not_found: List[OrderedDict] = []
    for pvo in pvos:
        rows = find_pvo(index, pvo)
        if rows:
            found_rows.append((pvo, rows))
        else:
            not_found.append(build_not_found_entry(pvo))
    return found_rows, not_found


def main() -> int:
    parser = argparse.ArgumentParser(description="Map Oracle BICC PVOs to BOSS Business Views using the Master sheet.")
    parser.add_argument("workbook", help="Path to Business_Object_Views_to_BICC_Database_Mapping_<version>.xlsx")
    parser.add_argument("--pvo", action="append", dest="pvos", help="A View Object / BICC PVO to map. Repeat for bulk use.")
    parser.add_argument("--pvo-file", help="Text file with one PVO per line for bulk use.")
    parser.add_argument("--mode", choices=["single", "summary", "columns"], default="summary", help="Execution mode: single, summary, or columns.")
    parser.add_argument("--csv-out", help="Write column-level mapping rows to CSV.")
    parser.add_argument("--pretty", action="store_true", help="Pretty-print JSON output.")
    args = parser.parse_args()

    if not args.pvos and not args.pvo_file:
        parser.error("Provide at least one --pvo or a --pvo-file")

    pvos: List[str] = []
    if args.pvos:
        pvos.extend(args.pvos)
    if args.pvo_file:
        with open(args.pvo_file, "r", encoding="utf-8") as f:
            pvos.extend([line.strip() for line in f if line.strip()])

    if args.mode == "single" and len(pvos) != 1:
        parser.error("mode=single requires exactly one PVO")

    records = load_rows(args.workbook)
    index = build_index(records)
    version = extract_version(args.workbook)
    found_rows, not_found = collect_results(pvos, index)

    payload: OrderedDict
    if args.mode == "single":
        if not found_rows:
            raise ValueError("Requested PVO is not present in workbook")
        pvo, rows = found_rows[0]
        payload = summarize_single(rows, pvo, version)
    elif args.mode == "summary":
        payload = OrderedDict([
            ("version", version),
            ("requested_pvos", pvos),
            ("found_count", len(found_rows)),
            ("not_found_count", len(not_found)),
            ("not_found", not_found),
        ])
    else:
        csv_target = args.csv_out if args.csv_out else default_csv_out_path(version)
        csv_written_to = resolve_csv_out_path(csv_target, version)
        csv_dir = os.path.dirname(csv_written_to)
        if csv_dir:
            os.makedirs(csv_dir, exist_ok=True)

        mappings = [
            summarize_mapping_with_columns(rows, pvo, version)
            for pvo, rows in found_rows
        ]
        write_csv(mappings, csv_written_to)
        row_count = sum(len(mapping.get("columns", [])) for mapping in mappings)
        payload = OrderedDict([
            ("version", version),
            ("csv_path", csv_written_to),
            ("row_count", row_count),
            ("found_count", len(found_rows)),
            ("not_found", not_found),
        ])

    json.dump(payload, sys.stdout, indent=2 if args.pretty else None)
    sys.stdout.write("\n")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
